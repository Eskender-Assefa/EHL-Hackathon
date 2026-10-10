from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .models import RequirementCandidate


def export_deck(requirements: list[RequirementCandidate], destination: str | Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "BMW Product Requirements Priorities"
    title_slide.placeholders[1].text = "Evidence-linked PM review deck\nGenerated from the automotive data pipeline"
    summary = presentation.slides.add_slide(presentation.slide_layouts[5])
    summary.shapes.title.text = "Decision overview"
    table = summary.shapes.add_table(min(len(requirements), 10) + 1, 6, Inches(0.4), Inches(1.3), Inches(12.5), Inches(5.5)).table
    headers = ["Band", "Platform", "Category", "Requirement", "Sentiment", "EHI"]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for row, requirement in enumerate(requirements[:10], start=1):
        values = [requirement.review_band, requirement.target_platform, requirement.taxonomy_category,
                  requirement.requirement_statement, f"{requirement.average_sentiment:+.0f}", f"{requirement.ehi_score:.0f}%"]
        for column, value in enumerate(values):
            table.cell(row, column).text = str(value)
    for requirement in requirements[:20]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"{requirement.review_band}: {requirement.title}"
        evidence = "\n".join(f"• {item.source_uri} — {item.locator}" for item in requirement.evidence[:4]) or "• No evidence"
        blindspots = "\n".join(f"• {item}" for item in requirement.blindspots) or "• None recorded"
        slide.placeholders[1].text = (
            f"{requirement.requirement_statement}\n\n"
            f"Category: {requirement.taxonomy_category} | Region: {requirement.region_variant}\n"
            f"Sentiment: {requirement.average_sentiment:+.1f} | EHI: {requirement.ehi_score:.1f}% | "
            f"Priority: {requirement.priority_score if requirement.priority_score is not None else 'Unscored'}\n\n"
            f"Evidence\n{evidence}\n\nBlindspots\n{blindspots}"
        )
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(16)
    presentation.save(destination)
    return destination
