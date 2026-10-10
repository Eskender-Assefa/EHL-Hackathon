import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from bmw_signals.bmw_topics import GENERIC_TOPICS, topic_category
from bmw_signals.local_ingestion import ingest_feedback_workbook
from bmw_signals.models import EvidenceRef, FilteredSignal, RawPost
from bmw_signals.pipeline import UNCLASSIFIED_TOPIC, RuleBasedEvaluator, _category, _sentiment, process
from bmw_signals.requirements import (
    MIN_REGION_SAMPLE, MIN_RESPONDENTS, MIN_SEVERE_SAFETY, build_requirements, summarise_exclusions,
)

STRICT = dict(min_respondents=MIN_RESPONDENTS, min_region_sample=MIN_REGION_SAMPLE, min_severe_safety=MIN_SEVERE_SAFETY)
MATERIAL = Path(__file__).resolve().parents[2] / "material"
HEADER = ["ID", "Source", "Brand", "Sales Description", "Country", "Derivate (e-code)", "Engine Type", "Feedback Type",
          "Customer Feedback", "Vfc level2 3 Name", "Vfc level2 Name", "Vfc level3 Name"]


def workbook_post(body, labels, record_id="G60-1"):
    return RawPost(
        source_platform="BMW_Workbook", source_url="bmw-file://G60_feedback_hackathon.xlsx", thread_title="G60",
        post_body=body, target_platform="G60", collected_at="2026-10-10T10:00:00Z", source_record_id=record_id,
        region="US", feedback_types=sorted({kind for _, kind in labels}),
        bmw_labels=[dict(topic=topic, feedback_type=kind, row_number=index) for index, (topic, kind) in enumerate(labels, start=2)],
    )


def signal(identifier, *, kinds=("Defect",), sentiment=-80, region="US", category="C07", topic="Touch screen, operation"):
    return FilteredSignal(
        signal_id=identifier, source_platform="BMW_Workbook", source_url="bmw-file://feedback.xlsx",
        target_platform="G60", taxonomy_category=category, cleaned_verbatim_text=f"Evidence {identifier}",
        sentiment_score=sentiment, confidence=0.9, topic=topic, region=region, source_record_id=identifier,
        feedback_types=list(kinds), classification_basis="bmw_label",
        evidence=[EvidenceRef(source_uri="bmw-file://feedback.xlsx", locator=f"row {identifier}", record_id=identifier)],
    )


class IngestionLabelTests(unittest.TestCase):
    def test_workbook_rows_keep_topic_and_type_paired(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "G60_feedback_hackathon.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Feedback_Explorer"
            sheet.append(HEADER)
            sheet.append(["G60-1", "A", "BMW", "5", "CA", "G60", "BEV", "Likes", "Seats are good but touchscreen is slow.", "x", "Seats", "Like"])
            sheet.append(["G60-1", "A", "BMW", "5", "CA", "G60", "BEV", "Difficult to Use", "Seats are good but touchscreen is slow.", "x", "Touch screen, operation", "Not friendly"])
            sheet.append(["G60-2", "A", "BMW", "5", "US", "G60", "ICE", None, "Great car.", "no_class_found", "no_class_found", None])
            workbook.save(path)
            posts, _ = ingest_feedback_workbook(path)
        pairs = [(label.topic, label.feedback_type, label.row_number) for label in posts[0].bmw_labels]
        self.assertEqual(pairs, [("Seats", "Likes", 2), ("Touch screen, operation", "Difficult to Use", 3)])
        self.assertEqual(posts[1].bmw_labels, [])
        self.assertEqual(posts[0].region, "RoW")


class LabelEvaluationTests(unittest.IsolatedAsyncioTestCase):
    async def test_bmw_feedback_type_decides_the_sign_not_the_wording(self):
        raw = workbook_post("Customer would like a great, good, excellent bigger screen.", [("Touch screen, operation", "Wants")])
        signals, _ = await process([raw], RuleBasedEvaluator())
        self.assertEqual(len(signals), 1)
        self.assertLess(signals[0].sentiment_score, 0)
        self.assertEqual(signals[0].topic, "Touch screen, operation")
        self.assertEqual(signals[0].taxonomy_category, "C07")
        self.assertEqual(signals[0].classification_basis, "bmw_label")
        self.assertEqual(signals[0].feedback_types, ["Wants"])

    async def test_one_signal_per_bmw_topic_even_within_one_category(self):
        raw = workbook_post("CarPlay drops and the navigation is confusing.", [("Apple CarPlay", "Defect"), ("Navigation", "Difficult to Use")])
        signals, _ = await process([raw], RuleBasedEvaluator())
        self.assertEqual(sorted(item.topic for item in signals), ["Apple CarPlay", "Navigation"])
        self.assertEqual({item.taxonomy_category for item in signals}, {"C07"})

    async def test_praise_and_complaint_on_one_topic_stay_visible(self):
        raw = workbook_post("Seats look fine. The seat squeaks.", [("Seats", "Likes"), ("Seats", "Defect")])
        signals, _ = await process([raw], RuleBasedEvaluator())
        self.assertEqual(len(signals), 1)
        self.assertEqual(signals[0].feedback_types, ["Defect", "Likes"])

    async def test_unlabelled_text_without_keywords_is_unclassified(self):
        raw = workbook_post("Really love this car!", [])
        signals, _ = await process([raw], RuleBasedEvaluator())
        self.assertEqual([item.topic for item in signals], [UNCLASSIFIED_TOPIC])
        self.assertEqual(signals[0].classification_basis, "keyword_rule")


class KeywordRuleTests(unittest.TestCase):
    def test_keywords_match_whole_words_only(self):
        self.assertEqual(_category("Customer is happy with the total range and the appointment")[0], "C02")
        self.assertEqual(_category("apparently the rotation was approved")[1], UNCLASSIFIED_TOPIC)
        self.assertEqual(_category("The app keeps crashing")[0], "C07")
        self.assertEqual(_category("Both seats squeak")[0], "C05")

    def test_request_phrases_are_not_praise(self):
        self.assertLess(_sentiment("Customer would like a larger display", []), 0)
        self.assertLess(_sentiment("I wish it had ventilated seats", []), 0)
        self.assertGreater(_sentiment("I like the display", []), 0)

    def test_single_polarity_label_overrides_keywords(self):
        self.assertLess(_sentiment("It is great and good", ["Defect"]), 0)
        self.assertGreater(_sentiment("No problem or issue at all", ["Likes"]), 0)


class TopicMapTests(unittest.TestCase):
    def test_known_topics_and_normalisation(self):
        self.assertEqual(topic_category("Touch screen, operation"), "C07")
        self.assertEqual(topic_category("ELECTRIC  range"), "C02")
        self.assertEqual(topic_category("Seats"), "C05")
        self.assertIsNone(topic_category("A topic BMW has not used"))
        self.assertIn("Body equipment, overall technology", GENERIC_TOPICS)

    @unittest.skipUnless(MATERIAL.exists(), "BMW material folder not present")
    def test_every_topic_in_the_supplied_workbooks_is_mapped(self):
        unmapped = set()
        for path in sorted(MATERIAL.glob("*_feedback_hackathon.xlsx")):
            sheet = load_workbook(path, read_only=True, data_only=True)["Feedback_Explorer"]
            rows = sheet.iter_rows(values_only=True)
            index = [str(cell).strip() for cell in next(rows)].index("Vfc level2 Name")
            for values in rows:
                topic = values[index]
                if topic and topic != "no_class_found" and topic_category(str(topic)) is None:
                    unmapped.add(str(topic))
        self.assertEqual(unmapped, set())


class RequirementThresholdTests(unittest.TestCase):
    def test_minimum_respondents_for_observed_feedback(self):
        below = [signal(f"A{index}") for index in range(MIN_RESPONDENTS - 1)]
        self.assertEqual(build_requirements(below, **STRICT)[0], [])
        requirements, _ = build_requirements(below + [signal("LAST")], **STRICT)
        self.assertEqual(len(requirements), 1)
        self.assertEqual(requirements[0].distinct_signal_count, MIN_RESPONDENTS)

    def test_pain_and_praise_are_counted_per_respondent(self):
        signals = ([signal(f"P{index}", kinds=("Defect",)) for index in range(10)]
                   + [signal(f"L{index}", kinds=("Likes",), sentiment=75) for index in range(4)]
                   + [signal("BOTH", kinds=("Defect", "Likes"), sentiment=-2)])
        item = build_requirements(signals, **STRICT)[0][0]
        self.assertEqual((item.pain_respondents, item.praise_respondents), (11, 5))
        self.assertNotIn("preserve", item.requirement_statement.lower())

    def test_strength_is_worded_as_preserve(self):
        signals = [signal(f"L{index}", kinds=("Likes",), sentiment=75) for index in range(14)] + [signal("P", kinds=("Wants",), sentiment=-45)]
        item = build_requirements(signals, **STRICT)[0][0]
        self.assertIn("preserve", item.requirement_statement.lower())
        self.assertIn("Preserve", item.title)

    def test_unclassified_and_thin_topics_are_reported_not_proposed(self):
        signals = ([signal(f"U{index}", topic=UNCLASSIFIED_TOPIC, category="C10", kinds=(), sentiment=0) for index in range(40)]
                   + [signal(f"T{index}", topic="Cupholder", category="C08") for index in range(3)])
        requirements, _ = build_requirements(signals, **STRICT)
        self.assertEqual(requirements, [])
        report = summarise_exclusions(signals)
        self.assertEqual(report["unclassified_respondents"], 40)
        self.assertEqual(report["topics_below_minimum"], 1)
        self.assertEqual(report["respondents_in_topics_below_minimum"], 3)

    def test_safety_escalation_needs_repeated_severe_negative_evidence(self):
        def safety(identifier, **changes):
            return signal(identifier, category="C04", topic="Braking", **changes)
        one_defect = [safety("D1")] + [safety(f"L{index}", kinds=("Likes",), sentiment=75) for index in range(20)]
        self.assertNotEqual(build_requirements(one_defect, **STRICT)[0][0].review_band, "Critical review")
        repeated = [safety(f"D{index}") for index in range(3)]
        item = build_requirements(repeated, **STRICT)[0][0]
        self.assertEqual(item.review_band, "Critical review")
        self.assertTrue(any("Low evidence volume" in warning for warning in item.warnings))

    def test_regional_conflict_needs_a_real_sample_in_each_region(self):
        us = [signal(f"US{index}", region="US") for index in range(20)]
        thin_eu = [signal(f"EU{index}", region="EU", kinds=("Likes",), sentiment=75) for index in range(3)]
        requirements, _ = build_requirements(us + thin_eu, **STRICT)
        self.assertEqual([item.region_variant for item in requirements], ["Global / Unassigned"])
        self.assertFalse(requirements[0].regional_conflict)
        full_eu = [signal(f"EU{index}", region="EU", kinds=("Likes",), sentiment=75) for index in range(15)]
        requirements, _ = build_requirements(us + full_eu, **STRICT)
        self.assertEqual({item.region_variant for item in requirements}, {"US", "EU"})

    def test_verified_evidence_is_reachable_with_bmw_labels(self):
        item = build_requirements([signal(f"A{index}") for index in range(60)], **STRICT)[0][0]
        self.assertEqual(item.confidence_label, "Verified Evidence")


if __name__ == "__main__":
    unittest.main()
