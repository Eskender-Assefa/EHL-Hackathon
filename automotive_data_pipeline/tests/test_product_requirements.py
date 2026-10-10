import json
import tempfile
import unittest
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation

from bmw_signals.audit import append_event, verify_audit
from bmw_signals.cli import _write_requirements
from bmw_signals.cockpit import create_app
from bmw_signals.deck import export_deck
from bmw_signals.local_ingestion import _pdf_publication_timestamp, ingest_feedback_workbook
from bmw_signals.models import EvidenceRef, FilteredSignal, ScoringInputs
from bmw_signals.pipeline import RuleBasedEvaluator, process
from bmw_signals.requirements import build_requirements
from test_pipeline import post


def signal(
    identifier, *, region="EU", sentiment=-70, timestamp=None, topic="Touchscreen",
    scope="variable_later", source="BMW_Workbook", source_url="bmw-file://feedback.xlsx",
):
    return FilteredSignal(
        signal_id=identifier, source_platform=source, source_url=source_url,
        timestamp=timestamp, target_platform="G60", taxonomy_category="C07",
        cleaned_verbatim_text=f"Evidence {identifier}", sentiment_score=sentiment,
        confidence=0.8, topic=topic, requirement_statement="Improve touchscreen usability for G60 customers.",
        scope_layer=scope, region=region, source_record_id=identifier,
        evidence=[EvidenceRef(source_uri=source_url, locator=f"record {identifier}", source_platform=source, record_id=identifier)],
    )


class LocalIngestionTests(unittest.TestCase):
    def test_feedback_ids_are_collapsed_with_all_row_evidence(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "G60_feedback_hackathon.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Feedback_Explorer"
            sheet.append(["ID", "Source", "Brand", "Sales Description", "Country", "Derivate (e-code)", "Engine Type", "Feedback Type", "Customer Feedback", "Vfc level2 3 Name", "Vfc level2 Name", "Vfc level3 Name"])
            sheet.append(["G60-1", "A", "BMW", "5", "US", "G60", "BEV", "Likes", "Seats are good but touchscreen is slow.", "Seats", "Seats", "Like"])
            sheet.append(["G60-1", "A", "BMW", "5", "US", "G60", "BEV", "Difficult to Use", "Seats are good but touchscreen is slow.", "Touch screen", "Touch screen", "Not friendly"])
            sheet.append(["G60-2", "A", "BMW", "5", "GB", "G60", "ICE", "Defect", "Software crashes every morning.", "Software", "Software", "Malfunction"])
            workbook.save(path)
            posts, manifest = ingest_feedback_workbook(path)
        self.assertEqual(len(posts), 2)
        self.assertEqual(manifest["tagged_rows"], 3)
        self.assertEqual(manifest["duplicates_collapsed"], 1)
        self.assertEqual(len(posts[0].evidence), 2)
        self.assertEqual(posts[0].feedback_types, ["Likes", "Difficult to Use"])
        self.assertEqual(posts[0].region, "US")
        self.assertEqual(posts[1].region, "EU")

    def test_optionlist_date_with_pdf_spacing(self):
        self.assertEqual(_pdf_publication_timestamp("Ausgabe 2, 22.06. 2026"), "2026-06-22T00:00:00+00:00")


class RuleEngineTests(unittest.IsolatedAsyncioTestCase):
    async def test_mixed_feedback_is_split_and_scored(self):
        raw = post(post_body="The seats are comfortable, but the touch screen menu is terrible.")
        signals, audit = await process([raw], RuleBasedEvaluator())
        self.assertEqual([item.taxonomy_category for item in signals], ["C05", "C07"])
        self.assertGreater(signals[0].sentiment_score, 0)
        self.assertLess(signals[1].sentiment_score, 0)
        self.assertEqual(audit[0]["status"], "accepted")

    async def test_software_and_interior_rules_do_not_spill_into_safety(self):
        evaluator = RuleBasedEvaluator()
        cases = {
            "The forward-collision warning configuration menu requires six touchscreen taps and often freezes while driving.": "C07",
            "The dashboard display placement blocks the driver's sightline and makes the center controls hard to reach.": "C08",
            "Forward-collision warning alerts trigger false positives during normal motorway driving and demand immediate correction.": "C04",
            "I am happy with the overall ownership experience and would recommend the vehicle to friends.": "C10",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                evaluation = await evaluator.evaluate(post(post_body=text))
                self.assertEqual(evaluation.aspects[0].taxonomy_category, expected)

    async def test_record_level_safety_tag_does_not_contaminate_split_fragments(self):
        raw = post(
            source_platform="BMW_Workbook",
            source_url="bmw-file://feedback.xlsx",
            timestamp=None,
            post_body="The overall ownership experience is reasonable. The touchscreen menu freezes during every morning commute.",
            original_tags=["Blind Spot"],
        )
        evaluation = await RuleBasedEvaluator().evaluate(raw)
        self.assertEqual([aspect.taxonomy_category for aspect in evaluation.aspects], ["C10", "C07"])

    async def test_duplicate_scope_includes_platform_and_valid_record_wins(self):
        evaluator = RuleBasedEvaluator()
        valid = post(source_record_id="same")
        other_platform = post(source_record_id="same", target_platform="G70")
        signals, audit = await process([post(source_record_id="same", target_platform=None), valid, other_platform], evaluator)
        self.assertEqual({item.target_platform for item in signals}, {"G60", "G70"})
        self.assertEqual([item["status"] for item in audit], ["quarantined", "accepted", "accepted"])


class RequirementTests(unittest.TestCase):
    def test_missing_inputs_use_labelled_static_evidence_proxy(self):
        requirements, baseline = build_requirements([signal("A"), signal("B")])
        self.assertFalse(baseline)
        self.assertNotEqual(requirements[0].review_band, "Unscored")
        self.assertIsNotNone(requirements[0].priority_score)
        self.assertEqual(requirements[0].priority_score_basis, "static_evidence_proxy")
        self.assertEqual(len(requirements[0].blindspots), 3)
        self.assertTrue(requirements[0].title.startswith("G60 (C07): "))
        self.assertNotIn("for G60 customers", requirements[0].requirement_statement)

    def test_unknown_region_is_normalized_for_presentation(self):
        requirements, _ = build_requirements([signal("A", region="Unknown"), signal("B", region="Unknown")])
        self.assertEqual(requirements[0].region_variant, "Global / Unassigned")

    def test_complete_inputs_produce_deterministic_priority(self):
        signals = [
            signal("A", timestamp="2026-09-01T00:00:00Z"),
            signal("B", timestamp="2026-09-15T00:00:00Z"),
            signal("C", timestamp="2026-10-01T00:00:00Z"),
        ]
        inputs = {"G60|C07|touchscreen": ScoringInputs(sales_causality_score=20, competitor_roi_score=10, effort_class="software_ux")}
        requirements, _ = build_requirements(signals, inputs)
        item = requirements[0]
        self.assertIsNotNone(item.velocity_multiplier)
        self.assertIsNotNone(item.priority_score)
        self.assertEqual(item.priority_score_basis, "complete_formula")
        self.assertEqual(item.effort_penalty, 0)
        self.assertGreater(item.priority_score, 30)

    def test_regional_conflict_creates_variants(self):
        signals = [signal("EU1", region="EU", sentiment=-90), signal("EU2", region="EU", sentiment=-80),
                   signal("CN1", region="CN", sentiment=70), signal("CN2", region="CN", sentiment=60)]
        requirements, _ = build_requirements(signals)
        self.assertEqual({item.region_variant for item in requirements}, {"EU", "CN"})
        self.assertTrue(all(item.regional_conflict for item in requirements))
        self.assertTrue(all(item.confidence_label == "Uncertain / Conflict Flag" for item in requirements))

    def test_fixed_baseline_does_not_become_customer_requirement(self):
        requirements, baseline = build_requirements([signal("P1", scope="fixed_now")])
        self.assertFalse(requirements)
        self.assertEqual(len(baseline), 1)

    def test_multi_source_evidence_is_balanced_and_competitor_does_not_dilute_sentiment(self):
        signals = [
            signal("shared", sentiment=-70),
            signal("shared", sentiment=-80, source="Reddit", source_url="https://www.reddit.com/r/BMW/comments/abc"),
            signal("BP1", sentiment=-60, source="BIMMERPOST", source_url="https://g60.bimmerpost.com/forums/showthread.php?t=1"),
            signal("GCF1", sentiment=-75, source="German_Car_Forum", source_url="https://www.germancarforum.com/threads/bmw-g60.1/"),
            signal("BF1", sentiment=-65, source="Bimmerforums", source_url="https://www.bimmerforums.com/forum/showthread.php?1"),
            signal("MB1", sentiment=0, scope="hypothetical_later", source="Competitor_Web", source_url="https://www.mercedes-benz.com/en/innovation/software/"),
        ]
        requirements, _ = build_requirements(signals)
        self.assertEqual(len(requirements), 1)
        item = requirements[0]
        self.assertEqual(item.distinct_signal_count, 5)
        self.assertEqual(item.average_sentiment, -70)
        self.assertEqual(set(item.evidence_source_counts), {"BMW_Workbook", "Reddit", "BIMMERPOST", "German_Car_Forum", "Bimmerforums", "Competitor_Web"})
        self.assertEqual([evidence.source_platform for evidence in item.evidence[:6]], ["Reddit", "BMW_Workbook", "BIMMERPOST", "Competitor_Web", "German_Car_Forum", "Bimmerforums"])
        with tempfile.TemporaryDirectory() as temporary:
            destination = export_deck(requirements, Path(temporary) / "multi-source.pptx")
            presentation = Presentation(destination)
            text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("reddit.com", text)
        self.assertIn("bimmerpost.com", text)
        self.assertIn("germancarforum.com", text)
        self.assertIn("bimmerforums.com", text)
        self.assertIn("mercedes-benz.com", text)


class AuditCockpitDeckTests(unittest.TestCase):
    def test_audit_chain_detects_tampering(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "audit.jsonl"
            append_event(path, actor="AI", requirement_id="REQ1", action="Propose", field_changed="requirement", previous_value=None, new_value="x", justification="fixture")
            self.assertEqual(len(verify_audit(path)), 1)
            path.write_text(path.read_text(encoding="utf-8").replace('"new_value": "x"', '"new_value": "tampered"'), encoding="utf-8")
            with self.assertRaises(ValueError):
                verify_audit(path)

    def test_decisions_require_rationale_and_append_audit(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            requirements, _ = build_requirements([signal("A"), signal("B")])
            (directory / "requirements.json").write_text(json.dumps([item.model_dump() for item in requirements]), encoding="utf-8")
            client = TestClient(create_app(directory))
            identifier = requirements[0].requirement_id
            self.assertEqual(client.post(f"/api/requirements/{identifier}/decision", json={"actor": "PM", "action": "Reject", "justification": ""}).status_code, 400)
            response = client.post(f"/api/requirements/{identifier}/decision", json={"actor": "PM", "action": "Postpone", "justification": "Await dated evidence"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["requirement"]["status"], "Postponed")
            self.assertEqual(client.get("/api/audit/verify").json()["events"], 1)

    def test_powerpoint_export(self):
        with tempfile.TemporaryDirectory() as temporary:
            requirements, _ = build_requirements([signal("A"), signal("B")])
            requirements[0] = requirements[0].model_copy(update={"priority_score": None, "review_band": "Unscored"})
            destination = export_deck(requirements, Path(temporary) / "deck.pptx")
            self.assertTrue(destination.exists())
            presentation = Presentation(destination)
            self.assertGreaterEqual(len(presentation.slides), 3)
            self.assertAlmostEqual(presentation.slide_width / presentation.slide_height, 16 / 9, places=2)
            text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
            self.assertNotIn("Unscored", text)
            self.assertNotIn("Region: Unknown", text)
            self.assertIn("G60 (C07):", text)
            self.assertIn("Static evidence proxy", text)

    def test_powerpoint_reserves_a_detail_slide_for_external_evidence(self):
        requirements, _ = build_requirements([signal("baseA"), signal("baseB")])
        base = requirements[0]
        internal = [
            base.model_copy(update={"requirement_id": f"INTERNAL{i:02}", "title": f"Internal priority {i}", "priority_score": 90 - i})
            for i in range(21)
        ]
        competitor_uri = "https://media.byd.com/example-benchmark"
        competitor = base.model_copy(update={
            "requirement_id": "EXTERNAL01", "title": "Neue_Klasse (C02): Improve charging predictability",
            "priority_score": 1, "review_band": "Low", "evidence_source_counts": {"Competitor_Web": 1},
            "evidence": [EvidenceRef(source_uri=competitor_uri, locator="official release", source_platform="Competitor_Web")],
        })
        with tempfile.TemporaryDirectory() as temporary:
            destination = export_deck([*internal, competitor], Path(temporary) / "deck.pptx")
            presentation = Presentation(destination)
            text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("media.byd.com", text)

    def test_manual_trigger_creates_labelled_neue_klasse_scenario(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            (directory / "requirements.json").write_text("[]", encoding="utf-8")
            (directory / "signals.json").write_text("[]", encoding="utf-8")
            client = TestClient(create_app(directory))
            response = client.post("/api/trigger-synthetic", json={"actor": "PM", "seed": 7, "justification": "Demo regional branching"})
            self.assertEqual(response.status_code, 200)
            self.assertIn("hypothetical", response.json()["warning"])
            generated = client.get("/api/requirements").json()
            self.assertTrue(generated)
            self.assertEqual({item["target_platform"] for item in generated}, {"Neue_Klasse"})
            self.assertTrue(all(any("synthetic" in warning.lower() for warning in item["warnings"]) for item in generated))

    def test_pipeline_recompute_preserves_pm_decision(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            signals = [signal("A"), signal("B")]
            generated = _write_requirements(directory, signals)
            client = TestClient(create_app(directory))
            response = client.post(f"/api/requirements/{generated[0].requirement_id}/decision", json={"actor": "PM", "action": "Approve", "justification": ""})
            self.assertEqual(response.status_code, 200)
            _write_requirements(directory, signals)
            saved = json.loads((directory / "requirements.json").read_text(encoding="utf-8"))
            self.assertEqual(saved[0]["status"], "Approved")
