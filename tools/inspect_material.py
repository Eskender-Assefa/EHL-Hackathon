"""Read-only source layout inventory. Never prints full customer statements."""
import json
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1]
for path in sorted((root/'material').glob('*feedback*.xlsx')):
    book = load_workbook(path,read_only=True,data_only=True)
    for sheet in book:
        rows=sheet.iter_rows(values_only=True)
        header=next(rows)
        first=next(rows)
        print(json.dumps({'file':path.name,'sheet':sheet.title,'rows':sheet.max_row,
                          'header':header,'sample_types':[type(x).__name__ for x in first]},default=str))
    book.close()
for path in sorted((root/'material').glob('*OptionList.*')):
    reader=PdfReader(path)
    print(json.dumps({'file':path.name,'pages':len(reader.pages),
                      'opening':reader.pages[0].extract_text()[:1500]},ensure_ascii=True))
