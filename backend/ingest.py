"""Explicit local source imports; no browser uploads, scraping or AI calls."""
import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import date,datetime
from pathlib import Path
from urllib.parse import urlsplit
from openpyxl import load_workbook
from pypdf import PdfReader
from psycopg.types.json import Jsonb
import bmwdb

HEADERS=['ID','Source','Brand','Sales Description','Country','Derivate (e-code)','Engine Type','Feedback Type','Customer Feedback','Vfc level2 3 Name','Vfc level2 Name','Vfc level3 Name']
COUNTRIES=set('DE GB US CN JP SE FI KR AU ZA FR DK NO CA LV ES'.split())
MAPPING_VERSION='baseline-v2.0-keywords-1 (provisional)'
RULES=[('C09',r'parking|camera|cruise|lane|driver.assist|automated|autonomous|collision.assist'),
       ('C07',r'carplay|idrive|software|navigation|bluetooth|touch.screen|menu|phone|update|digital|voice.control'),
       ('C03',r'price|cost|subscription|afford|expensive|value.for.money'),
       ('C08',r'upholster|leather|material|interior.design|trim|ambient.light|plastic'),
       ('C04',r'brake|tyre|tire|suspension|steer|airbag|stability|safety'),
       ('C05',r'seat|lumbar|climate|ventilat|hvac|air.condition|audio|noise|comfort|key|tailgate|access'),
       ('C02',r'range|charg|battery|emission|fuel|electricity|wltp|consumption'),
       ('C06',r'exterior|headlight|paint|glaz|roof|wheel|dimension|width|length|height'),
       ('C01',r'engine|drivetrain|transmission|power|torque|acceleration|speed|motor')]


def classify(tags):
    text=' '.join(str(x or '') for x in tags)
    for code,pattern in RULES:
        if re.search(pattern,text,re.I):return code
    return None


def archive(path):
    content=Path(path).read_bytes(); digest=hashlib.sha256(content).hexdigest()
    stored=Path(__file__).resolve().parents[1]/'data/assets'/digest
    stored.parent.mkdir(parents=True,exist_ok=True)
    try:
        with stored.open('xb') as f:f.write(content)
    except FileExistsError:
        if stored.read_bytes()!=content:raise ValueError('Archive hash mismatch')
    return digest,stored


def begin(conn,path,namespace,kind,origin,metadata,parser):
    digest,stored=archive(path)
    conn.execute('SELECT pg_advisory_xact_lock(709102027)')
    conn.execute('INSERT INTO datasets(namespace,evidence_kind,metadata) VALUES (%s,%s,%s) ON CONFLICT(namespace) DO NOTHING',(namespace,kind,Jsonb(metadata)))
    dataset,existing_kind=conn.execute('SELECT id,evidence_kind FROM datasets WHERE namespace=%s',(namespace,)).fetchone()
    if existing_kind!=kind:raise ValueError('Namespace evidence kind mismatch')
    conn.execute('INSERT INTO source_assets(sha256,original_filename,original_location,storage_path) VALUES (%s,%s,%s,%s) ON CONFLICT(sha256) DO NOTHING',(digest,Path(path).name,str(Path(path).resolve()),str(stored.resolve())))
    asset=conn.execute('SELECT id FROM source_assets WHERE sha256=%s',(digest,)).fetchone()[0]
    inserted=conn.execute('INSERT INTO dataset_versions(dataset_id,asset_id,data_origin,metadata) VALUES (%s,%s,%s,%s) ON CONFLICT(dataset_id,asset_id) DO NOTHING RETURNING id',(dataset,asset,origin,Jsonb(metadata))).fetchone()
    if not inserted:
        version,prior_origin=conn.execute('SELECT id,data_origin FROM dataset_versions WHERE dataset_id=%s AND asset_id=%s',(dataset,asset)).fetchone()
        if prior_origin!=origin:raise ValueError('Snapshot origin mismatch')
        return version,None
    version=inserted[0]
    run=conn.execute("INSERT INTO ingestion_runs(asset_id,version_id,parser_version,status) VALUES (%s,%s,%s,'running') RETURNING id",(asset,version,parser)).fetchone()[0]
    return version,run


def raw_record(conn,version,number,payload,sheet=None,page=None,external=None):
    record=conn.execute('INSERT INTO source_records(version_id,record_number,external_id,raw_payload,scope_layer) VALUES (%s,%s,%s,%s,%s) RETURNING id',(version,number,external,Jsonb(payload),'fixed')).fetchone()[0]
    conn.execute('INSERT INTO evidence_locations(record_id,page,sheet,row_number,cell_range) VALUES (%s,%s,%s,%s,%s)',(record,page,sheet,number if page is None else None,f'A{number}:L{number}' if sheet else None))
    return record


def issue(conn,run,record,reason):
    conn.execute('INSERT INTO validation_issues(run_id,record_id,code,reason) VALUES (%s,%s,%s,%s)',(run,record,'source_mapping',reason))


def finish(conn,run,report,has_issues=False):
    conn.execute('UPDATE ingestion_runs SET status=%s,finished_at=now(),report=%s WHERE id=%s',('quarantined' if has_issues else 'succeeded',Jsonb(report),run))


def import_feedback(conn,path,model):
    book=load_workbook(path,read_only=True,data_only=False)
    if 'Feedback_Explorer' not in book.sheetnames:raise ValueError('Feedback_Explorer sheet required')
    sheet=book['Feedback_Explorer']; iterator=sheet.iter_rows(values_only=True)
    if list(next(iterator))!=HEADERS:raise ValueError('Unexpected feedback header')
    rows=list(iterator);book.close()
    metadata={'model_code':model,'sheet':'Feedback_Explorer','publication_date':None,
              'origin_note':'Supplied hackathon workbook; real-world provenance unverified',
              'category_mapping_version':MAPPING_VERSION,'variant_note':'Source labels retained; canonical variant unresolved'}
    parents={};errors=0; tags=0
    with conn.transaction():
        version,run=begin(conn,path,f'bmw.feedback.{model}','customer_statement','unknown',metadata,'feedback-xlsx-1')
        if run is None:return {'version_id':str(version),'reimport':True}
        for number,values in enumerate(rows,2):
            payload={key:(value.isoformat() if isinstance(value,(datetime,date)) else value) for key,value in zip(HEADERS,values)}
            # Row identity, not parent feedback ID: repeated IDs are legitimate tags.
            record=raw_record(conn,version,number,payload,sheet=sheet.title,external=f'{sheet.title}!{number}')
            eid=payload['ID'];statement=payload['Customer Feedback'];country=payload['Country']
            if not eid or not isinstance(statement,str) or not statement.strip() or payload['Derivate (e-code)']!=model:
                issue(conn,run,record,'Missing feedback identity/text or model mismatch');errors+=1;continue
            market=country if country in COUNTRIES else None
            if market:conn.execute("INSERT INTO markets VALUES (%s,'country') ON CONFLICT DO NOTHING",(market,))
            if eid in parents:
                parent,previous_text,previous_country=parents[eid]
                if statement!=previous_text or country!=previous_country:
                    issue(conn,run,record,'Conflicting content/context for repeated feedback ID; raw row retained');errors+=1;continue
            else:
                parent=conn.execute('''INSERT INTO feedback_records(version_id,external_id,model_code,market_code,raw_market,statement,source_platform,feedback_type)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id''',(version,str(eid),model,market,country,statement,payload['Source'],payload['Feedback Type'])).fetchone()[0]
                parents[eid]=(parent,statement,country)
                category=classify(values[9:12])
                conn.execute('''INSERT INTO evidence_items(version_id,source_record_id,feedback_id,model_code,market_code,category_id,category_mapping_version,evidence_kind,source_platform,title,statement,data_origin,safety_review,context)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,'customer_statement',%s,%s,%s,'unknown',%s,%s)''',
                    (version,record,parent,model,market,category,MAPPING_VERSION,payload['Source'],statement[:150],statement,
                     bool(re.search(r'brak|airbag|fire|crash|injur|danger|safety|collision',statement,re.I)),
                     Jsonb({'external_feedback_id':str(eid),'raw_market':country,'engine_type':payload['Engine Type'],
                            'sales_description':payload['Sales Description'],'feedback_type':payload['Feedback Type'],
                            'mapping_note':'Provisional keyword mapping from first tag row; original tags preserved',
                            'variant_applicability':'unknown'})))
            original={key:payload[key] for key in HEADERS[7:] if key!='Customer Feedback'}
            conn.execute('INSERT INTO feedback_tags(feedback_id,version_id,source_record_id,original_tags) VALUES (%s,%s,%s,%s)',(parent,version,record,Jsonb(original)));tags+=1
        report={'version_id':str(version),'raw_rows':len(rows),'feedback_records':len(parents),'tag_rows':tags,'quarantined_rows':errors,'data_origin':'unknown'}
        finish(conn,run,report,errors>0)
        return report


def import_catalogue(conn,path,model):
    reader=PdfReader(path)
    if reader.is_encrypted:reader.decrypt('')
    pages=[p.extract_text() or '' for p in reader.pages]
    opening=pages[0]
    # Retain exact edition wording rather than infer a date or variant globally.
    edition_lines=[x.strip() for x in opening.splitlines() if re.search(r'\b(2026|Stand|Ausgabe|Issue)\b',x,re.I)]
    metadata={'model_code':model,'edition_text':' | '.join(edition_lines) or None,
              'market_applicability':'unknown','variant_applicability':'unknown',
              'extraction_note':'Page-level text context, not reviewed structured availability rules',
              'language_context':'German-language catalogue; no universal market applicability'}
    with conn.transaction():
        version,run=begin(conn,path,f'bmw.catalogue.{model}','published_specification','observed',metadata,'catalogue-page-extract-1')
        if run is None:return {'version_id':str(version),'reimport':True}
        extracted=0
        for page,text in enumerate(pages,1):
            record=raw_record(conn,version,page,{'page':page,'text':text},page=page,external=f'page:{page}')
            if not text.strip():issue(conn,run,record,'No extractable page text; OCR not performed');continue
            conn.execute('''INSERT INTO evidence_items(version_id,source_record_id,model_code,evidence_kind,title,statement,data_origin,context,mapping_status)
                VALUES (%s,%s,%s,'published_specification',%s,%s,'observed',%s,'unmapped')''',
                (version,record,model,f'{model} catalogue - page {page}',text,Jsonb(metadata)));extracted+=1
        report={'version_id':str(version),'pages':len(pages),'page_extracts':extracted,'review_status':'unreviewed','market_applicability':'unknown'}
        finish(conn,run,report,extracted!=len(pages));return report


def import_pipeline(conn,path):
    """Fixture/output adapter: raw records + exact signal excerpts + run metadata.

    No body-only deduplication and no missing-date suppression. Each source/context
    is preserved. Pipeline C04/C05 signals remain ambiguous pending taxonomy review.
    """
    from automotive_data_pipeline.bmw_signals.models import RawPost,FilteredSignal
    bundle=json.loads(Path(path).read_text(encoding='utf-8'))
    raw=[RawPost.model_validate(x) for x in bundle['raw']]
    signals=[FilteredSignal.model_validate(x) for x in bundle['signals']]
    run_meta=bundle['run']
    for key in ['algorithm_version','model_name','prompt_version','taxonomy_version']:
        if key not in run_meta:raise ValueError('Pipeline run metadata required: '+key)
    origin=bundle.get('data_origin','unknown')
    if origin not in ['observed','synthetic','unknown']:raise ValueError('Invalid pipeline origin')
    with conn.transaction():
        version,run=begin(conn,path,'pipeline.'+bundle['namespace'],'derived_assessment',origin,run_meta,'pipeline-adapter-1')
        if run is None:return {'version_id':str(version),'reimport':True}
        errors=0;accepted=0
        for number,post in enumerate(raw,1):
            record=raw_record(conn,version,number,post.model_dump(),external=str(number))
            url=urlsplit(post.source_url)
            if url.scheme not in ['http','https'] or not url.hostname or post.target_platform not in ['F70','G60','G70']:
                issue(conn,run,record,'Invalid source URL or unsupported model context');errors+=1;continue
            matches=[s for s in signals if s.source_url==post.source_url and s.target_platform==post.target_platform and s.source_platform==post.source_platform and s.cleaned_verbatim_text in post.post_body and not s.is_noise]
            # One evidence statement retains full text plus all extracted aspects.
            category=matches[0].taxonomy_category if len({s.taxonomy_category for s in matches})==1 else None
            ambiguous=category in ['C04','C05']
            comparison_context=bundle.get('contexts',{}).get(post.source_url,{}) or {}
            market=comparison_context.get('market_code')
            if market not in COUNTRIES:market=None
            if market:conn.execute("INSERT INTO markets VALUES (%s,'country') ON CONFLICT DO NOTHING",(market,))
            published=datetime.fromisoformat(post.timestamp.replace('Z','+00:00')).date() if post.timestamp else None
            context={'original':post.model_dump(),'pipeline_aspects':[s.model_dump() for s in matches],
                     'benchmark_context':post.benchmark_context,'source_manufacturer':comparison_context.get('source_manufacturer','unknown'),
                     'comparison_context':comparison_context,
                     'taxonomy_note':'C04/C05 definitions differ; human mapping required' if ambiguous else 'Original pipeline category preserved, provisional',
                     'run_metadata':run_meta}
            conn.execute('''INSERT INTO evidence_items(version_id,source_record_id,model_code,market_code,category_id,category_mapping_version,evidence_kind,source_platform,title,statement,source_url,publication_date,captured_at,data_origin,context,safety_review)
                VALUES (%s,%s,%s,%s,%s,%s,'derived_assessment',%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
                (version,record,post.target_platform,market,None if ambiguous else category,'pipeline-to-baseline-1',post.source_platform,
                 post.thread_title or 'Collected source',post.post_body,post.source_url,published,post.collected_at,origin,Jsonb(context),ambiguous or bool(re.search(r'brak|fire|injur|safety',post.post_body,re.I))));accepted+=1
        conn.execute('''INSERT INTO processing_runs(version_id,algorithm_version,model_name,prompt_version,taxonomy_version,input_assumptions,status)
            VALUES (%s,%s,%s,%s,%s,%s,%s)''',(version,run_meta['algorithm_version'],run_meta['model_name'],run_meta['prompt_version'],run_meta['taxonomy_version'],Jsonb(run_meta),'quarantined' if errors else 'succeeded'))
        report={'version_id':str(version),'raw_records':len(raw),'evidence_records':accepted,'quarantined_rows':errors}
        finish(conn,run,report,errors>0);return report


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('kind',choices=['feedback','catalogue','pipeline','material'])
    parser.add_argument('path',nargs='?',default='material')
    parser.add_argument('--model',choices=['F70','G60','G70'])
    args=parser.parse_args()
    with bmwdb.connect() as conn:
        if args.kind=='material':
            for model in ['F70','G60','G70']:
                for kind,filename,fn in [('feedback',f'{model}_feedback_hackathon.xlsx',import_feedback),('catalogue',f'{model}_OptionList.PDF',import_catalogue)]:
                    print(json.dumps({'source':filename,**fn(conn,Path(args.path)/filename,model)},indent=2))
        elif args.kind=='pipeline':print(json.dumps(import_pipeline(conn,args.path),indent=2))
        else:
            if not args.model:parser.error('--model required')
            fn=import_feedback if args.kind=='feedback' else import_catalogue
            print(json.dumps(fn(conn,args.path,args.model),indent=2))

if __name__=='__main__':main()
